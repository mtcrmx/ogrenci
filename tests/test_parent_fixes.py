"""Parent navigation, missing data and result ownership regressions."""
import unittest
from test_app_workflows import w, d


class ParentFixes(unittest.TestCase):
    def setUp(self):
        con = d._conn()
        self.students = [dict(r) for r in con.execute("SELECT o.* FROM ogrenciler o JOIN siniflar s ON s.id=o.sinif_id WHERE s.sinif_adi='8/A' LIMIT 2")]
        con.close()
        self.client = w.app.test_client()
        with self.client.session_transaction() as s:
            s['veli_ogrenci_id'] = self.students[0]['id']

    def test_current_parent_login_still_works(self):
        client = w.app.test_client()
        response = client.post('/veli/giris', data={'ogr_no':self.students[0]['ogr_no'], 'sifre':w.VELI_SIFRE})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(client.get('/veli').status_code, 200)

    def test_navigation_and_missing_data(self):
        hub = self.client.get('/veli/analizler').get_data(as_text=True)
        for target in ('/evde-analiz', '/veli/deneme-analiz'):
            self.assertIn(target, hub)
        self.assertNotIn('/okul-yaprak-analiz', hub)
        home = self.client.get('/veli').get_data(as_text=True)
        self.assertNotIn('Geçen ayla aynı tempoda.', home)
        self.assertNotIn('0 hafta tam ödev', home)
        analysis = self.client.get('/evde-analiz').get_data(as_text=True)
        self.assertIn('Veli menüsü', analysis)
        self.assertIn('Tüm analizler', analysis)
        self.assertNotIn('ödev tam 0', self.client.get('/veli/karne').get_data(as_text=True))

    def test_result_link_cannot_show_other_student(self):
        result = d.lgs_deneme_ekle(self.students[1]['id'], 'Private result fixture', '2026-10-17', [dict(ders='Matematik', soru_sayisi=10, dogru=5, yanlis=2, bos=3)], tur='ders')
        page = self.client.get('/veli/sonuclar', query_string={'sonuc':result['id']}).get_data(as_text=True)
        self.assertNotIn('Private result fixture', page)
        self.assertIn('Bu sonuç artık mevcut değil veya bu öğrenciye ait değil.', page)

    def test_new_notification_links_to_own_result(self):
        oid = self.students[0]['id']
        result = d.lgs_deneme_ekle(oid, 'Linked result fixture', '2026-10-17', [dict(ders='Matematik', soru_sayisi=10, dogru=5, yanlis=2, bos=3)], tur='ders')
        d.veli_haber_ekle(oid, 'Linked result saved.', sonuc_id=result['id'])
        page = self.client.get('/veli').get_data(as_text=True)
        self.assertIn('/veli/sonuclar?sonuc='+str(result['id']), page)
        self.assertIn('Linked result fixture', self.client.get('/veli/sonuclar', query_string={'sonuc':result['id']}).get_data(as_text=True))

    def test_export_distinguishes_missing_evaluation(self):
        from openpyxl import load_workbook
        row = dict(ad_soyad='Fixture', ogr_no=1, sinif_adi='5/A', okudu=0, okumadi=0, tam=0, eksik=0, yok=0, hedef_net=0, hedef='', denemeler='')
        sheet = load_workbook(w._karne_excel([row], 1)).active
        self.assertEqual(sheet['D2'].value, 'Henüz değerlendirilmedi')
        self.assertEqual(sheet['F2'].value, 'Henüz değerlendirilmedi')
        row['yok'] = 1
        sheet = load_workbook(w._karne_excel([row], 1)).active
        self.assertEqual(sheet['F2'].value, 0)
        self.assertEqual(sheet['H2'].value, 1)
