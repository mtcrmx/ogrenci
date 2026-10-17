"""Security boundaries exercised without the convenience browser test client."""
import os
import tempfile
import unittest
from pathlib import Path

sandbox = tempfile.TemporaryDirectory(prefix='security-tests-')
os.environ['OGR_TAKIP_DB_PATH'] = str(Path(sandbox.name) / 'test.db')
import web_app as w
import database as db
import account_security as auth
from flask.testing import FlaskClient
from werkzeug.security import check_password_hash


class Security(unittest.TestCase):
    def setUp(self):
        w.app.config['TESTING'] = True
        self.client = FlaskClient(w.app)
        con = db._conn()
        con.execute('DELETE FROM login_attempts')
        self.students = [dict(r) for r in con.execute("SELECT o.* FROM ogrenciler o JOIN siniflar s ON s.id=o.sinif_id WHERE s.sinif_adi='8/A' LIMIT 2")]
        self.admin = dict(con.execute("SELECT * FROM ogretmenler WHERE ad_soyad='METEHAN CÜCEN'").fetchone())
        self.teacher = dict(con.execute("SELECT * FROM ogretmenler WHERE ad_soyad NOT IN ('METEHAN CÜCEN','ADEM AKGÜL') LIMIT 1").fetchone())
        con.commit(); con.close()

    def csrf(self, client=None):
        client = client or self.client
        response = client.get('/veli/giris')
        self.assertIn('name="_csrf_token"', response.get_data(as_text=True))
        with client.session_transaction() as s:
            return s['request_csrf']

    def parent_login(self, student, code, client=None):
        client = client or self.client
        return client.post('/veli/giris', data={'ogr_no': student['ogr_no'], 'sifre': code, '_csrf_token': self.csrf(client)})

    def teacher_session(self, teacher):
        with self.client.session_transaction() as s:
            s.update(ogretmen_id=teacher['id'], ogretmen_adi=teacher['ad_soyad'], ogretmen_yetki='tam', teacher_auth_version=auth.teacher_version(teacher['id']))

    def test_personal_codes_do_not_cross_students_and_reset_revokes_session(self):
        first, other = self.students
        code = auth.issue_parent_code(first['id'])
        auth.issue_parent_code(other['id'])
        self.assertEqual(self.parent_login(other, code).status_code, 200)
        with self.client.session_transaction() as s:
            self.assertNotIn('veli_ogrenci_id', s)
        self.assertEqual(self.parent_login(first, code).status_code, 302)
        self.assertEqual(self.client.get('/veli').status_code, 200)
        auth.issue_parent_code(first['id'])
        self.assertEqual(self.client.get('/veli').location, '/veli/giris')
        self.assertEqual(self.parent_login(first, code).status_code, 200)

    def test_legacy_shared_password_and_old_session_rejected(self):
        student = self.students[0]
        self.assertEqual(self.parent_login(student, 'veli').status_code, 200)
        with self.client.session_transaction() as s:
            self.assertNotIn('veli_ogrenci_id', s)
            s['veli_ogrenci_id'] = student['id']
        self.assertEqual(self.client.get('/veli/analizler').location, '/veli/giris')

    def test_password_migration_is_idempotent_and_preserves_password(self):
        con = db._conn()
        con.execute('UPDATE ogretmenler SET sifre=? WHERE id=?', ('fixture-legacy-password', self.teacher['id']))
        auth.migrate(con)
        first = con.execute('SELECT sifre FROM ogretmenler WHERE id=?', (self.teacher['id'],)).fetchone()[0]
        self.assertNotEqual(first, 'fixture-legacy-password')
        self.assertTrue(check_password_hash(first, 'fixture-legacy-password'))
        auth.migrate(con)
        self.assertEqual(first, con.execute('SELECT sifre FROM ogretmenler WHERE id=?', (self.teacher['id'],)).fetchone()[0])
        con.close()
        self.assertTrue(db.ogretmen_dogrula(self.teacher['ad_soyad'], 'fixture-legacy-password'))
        self.assertTrue(all('sifre' not in r for r in db.tum_sifre_listesi()))

    def test_csrf_required_for_form_and_json(self):
        self.assertEqual(self.client.post('/veli/giris', data={'ogr_no':58,'sifre':'veli'}).status_code, 400)
        self.teacher_session(self.admin)
        self.assertEqual(self.client.post('/admin/sifreler', data={'student_id':self.students[0]['id']}).status_code, 400)
        self.assertEqual(self.client.post('/api/ogretmen/ogrenci/1/veli-not', json={'not':'test'}).status_code, 400)
        token = self.csrf()
        result = self.client.post('/admin/sifreler', data={'student_id':self.students[0]['id']}, headers={'X-CSRF-Token':token})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.headers['Cache-Control'], 'no-store')

    def test_admin_uses_database_identity_and_teacher_reset_revokes_session(self):
        self.teacher_session(self.teacher)
        with self.client.session_transaction() as s:
            s['ogretmen_adi'] = 'METEHAN CÜCEN'
        self.assertEqual(self.client.get('/admin/sifreler').status_code, 403)
        auth.issue_teacher_code(self.teacher['id'])
        self.assertEqual(self.client.get('/').location, '/login')

    def test_login_rate_limit(self):
        for _ in range(10):
            self.assertEqual(self.parent_login(self.students[0], 'invalid').status_code, 200)
        self.assertEqual(self.parent_login(self.students[0], 'invalid').status_code, 429)

    def test_phone_token_reset_and_legacy_token_rejection(self):
        import veli_cihaz as device
        student = self.students[0]
        auth.issue_parent_code(student['id'])
        device.veli_cihaz_kaydet(student['id'], 'fixture-phone-token')
        self.assertEqual(device.veli_cihaz_ogrenci('fixture-phone-token'), student['id'])
        auth.issue_parent_code(student['id'])
        self.assertIsNone(device.veli_cihaz_ogrenci('fixture-phone-token'))
        con = db._conn()
        device._hazirla(con)
        import hashlib
        con.execute("INSERT OR REPLACE INTO veli_cihaz(token_hash,ogrenci_id,zaman) VALUES(?,?,datetime('now','+1 day'))", (hashlib.sha256(b'legacy-token').hexdigest(), student['id']))
        con.commit(); con.close()
        self.assertIsNone(device.veli_cihaz_ogrenci('legacy-token'))

    def test_result_link_is_owned_and_deleted_result_is_explained(self):
        student, other = self.students
        result = db.lgs_deneme_ekle(other['id'], 'Private result', '2026-10-17', [dict(ders='Matematik', soru_sayisi=10, dogru=5, yanlis=2, bos=3)], tur='ders')
        code = auth.issue_parent_code(student['id'])
        self.parent_login(student, code)
        page = self.client.get('/veli/sonuclar', query_string={'sonuc':result['id']}).get_data(as_text=True)
        self.assertNotIn('Private result', page)
        self.assertIn('Bu sonuç artık mevcut değil veya bu öğrenciye ait değil.', page)

    def test_export_distinguishes_missing_evaluation(self):
        from openpyxl import load_workbook
        row = dict(ad_soyad='Fixture', ogr_no=1, sinif_adi='5/A', okudu=0, okumadi=0, tam=0, eksik=0, yok=0, hedef_net=0, hedef='', denemeler='')
        sheet = load_workbook(w._karne_excel([row], 1)).active
        self.assertEqual(sheet['D2'].value, 'Henüz değerlendirilmedi')
        self.assertEqual(sheet['F2'].value, 'Henüz değerlendirilmedi')
        row['yok'] = 1
        sheet = load_workbook(w._karne_excel([row], 1)).active
        self.assertEqual(sheet['F2'].value, 0)
        self.assertEqual(sheet['H2'].value, 1)

    def test_parent_navigation_and_empty_states(self):
        student = self.students[0]
        code = auth.issue_parent_code(student['id'])
        self.parent_login(student, code)
        hub = self.client.get('/veli/analizler').get_data(as_text=True)
        for target in ('/evde-analiz', '/okul-yaprak-analiz', '/veli/deneme-analiz'):
            self.assertIn(target, hub)
        home = self.client.get('/veli').get_data(as_text=True)
        self.assertNotIn('Geçen ayla aynı tempoda.', home)
        self.assertNotIn('0 hafta tam ödev', home)
        analysis = self.client.get('/evde-analiz').get_data(as_text=True)
        self.assertIn('Veli menüsü', analysis)
        self.assertIn('Tüm analizler', analysis)
        report = self.client.get('/veli/karne').get_data(as_text=True)
        self.assertNotIn('ödev tam 0', report)


if __name__ == '__main__':
    unittest.main()
