import json
from io import BytesIO
import unittest
from unittest.mock import patch
import numpy as np
import cv2
import pypdfium2 as pdfium
from pypdf import PdfReader
from reportlab.pdfgen import canvas
from docx import Document
from test_app_workflows import w,d
import kamera_media as media
import kamera_test as kt
import lgs_optik as optik


def pdf(text='Question document'):
    out=BytesIO();c=canvas.Canvas(out);c.drawString(60,750,text);c.save();return out.getvalue()


def marked_form(exam,student,token,marks,extra=None,choices='ABCD'):
    data=media.form_pdf(exam,student,token,choices)
    with pdfium.PdfDocument(data) as doc:
        im=np.array(doc[0].render(scale=3).to_pil().convert('RGB'))
    for i,a in enumerate(marks):
        if a:
            x,y=media.bubble_position(i,choices.index(a))
            cv2.circle(im,(round(x*im.shape[1]/210),round(y*im.shape[0]/297)),round(1.65*im.shape[1]/210),(0,0,0),-1)
    if extra is not None:
        x,y=media.bubble_position(0,extra)
        cv2.circle(im,(round(x*im.shape[1]/210),round(y*im.shape[0]/297)),round(1.65*im.shape[1]/210),(0,0,0),-1)
    return im


class KeyParsing(unittest.TestCase):
    def test_table_labels_vertical_cells_and_unicode_separators(self):
        for text in ['Soru No: 1 2 3 4\nCevap: A B D C',
                     '1\n2\n3\n4\nA\nB\nD\nC',
                     '1\nA\n2\nB\n3\nD\n4\nC',
                     '1–A 2—B 3:(D) 4-C']:
            self.assertEqual(media.parse_key(text,4,'ABCD')[0],list('ABDC'),text)

    def test_pdf_footer_does_not_skip_ocr(self):
        with patch.object(media,'ocr_image',return_value='1 A 2 B 3 C 4 D') as ocr:
            text,used=media.extract_key_text(pdf('2026 - Page 1'),'.pdf')
            self.assertTrue(used)
            ocr.assert_called_once()
            self.assertEqual(media.parse_key(text,4,'ABCD')[0],list('ABCD'))

    def test_numbered_table_missing_duplicate_invalid(self):
        self.assertEqual(media.parse_key('CEVAP ANAHTARI\n1-C 2.A 3: D 4 B',4,'ABCD')[0],list('CADB'))
        self.assertEqual(media.parse_key('1 2 3 4\nA B D C',4,'ABCD')[0],list('ABDC'))
        a,warnings=media.parse_key('1 A\n1 B\n2 E\n4 C',3,'ABCD')
        self.assertEqual(a,['','','']);self.assertTrue(warnings)
        self.assertEqual(media.parse_key('ABC eğitim 2026',3,'ABCD')[0],['','',''])
        self.assertEqual(media.parse_key('1 A/B 2 C',2,'ABCD')[0],['','C'])

    def test_docx_key_keeps_cells_once(self):
        doc=Document();doc.add_paragraph('Cevap anahtarı');table=doc.add_table(rows=2,cols=3)
        for i in range(3):table.cell(0,i).text=str(i+1);table.cell(1,i).text='ABC'[i]
        out=BytesIO();doc.save(out)
        text,_=media.extract_key_text(out.getvalue(),'.docx')
        self.assertEqual(media.parse_key(text,3,'ABCD')[0],list('ABC'))

    def test_text_pdf_key(self):
        text,ocr=media.extract_key_text(pdf('1 A 2 B 3 C 4 D'),'.pdf')
        self.assertFalse(ocr);self.assertEqual(media.parse_key(text,4,'ABCD')[0],list('ABCD'))

    def test_word_converter_failure_is_explicit(self):
        doc=Document();doc.add_paragraph('Soru');out=BytesIO();doc.save(out)
        with patch('broadcast_presentations.office_command',return_value=None):
            with self.assertRaisesRegex(ValueError,'dönüştürücü'):media.convert_document(out.getvalue(),'.docx')

    def test_ninety_questions_five_choices(self):
        exam=dict(ad='Ninety question form',soru_sayisi=90)
        student=dict(ad_soyad='Example Student',sinif_adi='8/A',ogr_no=100)
        marks=list('ABCDE'*18);marks[7]=''
        im=marked_form(exam,student,'test-token-90',marks,choices='ABCDE')
        _,jpg=cv2.imencode('.jpg',im)
        result=media.read_form(jpg.tobytes(),90,'ABCDE')
        self.assertEqual(result['answers'],marks);self.assertEqual(result['uncertain'],[])


class CompressedPhoto(unittest.TestCase):
    def test_upload_size_rotation_and_shadow(self):
        exam=dict(ad='Upload test',soru_sayisi=90)
        student=dict(ad_soyad='Example Student',sinif_adi='8/A',ogr_no=100)
        marks=list('ABCDE'*18);marks[7]=''
        im=marked_form(exam,student,'compressed-photo',marks,choices='ABCDE')
        scale=2200/max(im.shape[:2])
        im=cv2.resize(im,None,fx=scale,fy=scale,interpolation=cv2.INTER_AREA)
        shade=np.linspace(.65,1,im.shape[1])[None,:,None]
        im=np.uint8(im*shade)
        im=cv2.rotate(im,cv2.ROTATE_90_CLOCKWISE)
        _,jpg=cv2.imencode('.jpg',im,[cv2.IMWRITE_JPEG_QUALITY,90])
        result=media.read_form(jpg.tobytes(),90,'ABCDE')
        self.assertEqual(result['answers'],marks)
        self.assertEqual(result['uncertain'],[])


class CameraWorkflow(unittest.TestCase):
    def setUp(self):
        self.t=next(t for t in d.tum_ogretmenler() if t['yetki']=='tam')
        self.c=w.app.test_client()
        with self.c.session_transaction() as s:s.update(ogretmen_id=self.t['id'],ogretmen_adi=self.t['ad_soyad'],ogretmen_yetki='tam')
        con=d._conn();self.student=dict(con.execute("SELECT o.*,s.sinif_adi FROM ogrenciler o JOIN siniflar s ON s.id=o.sinif_id WHERE s.sinif_adi='8/A' LIMIT 1").fetchone());con.close()
        self.assertEqual(self.c.get('/yaprak-kamera').status_code,200)
        with self.c.session_transaction() as s:self.csrf=s['kamera_csrf']
        self.before_ids={e['id'] for e in optik.sinavlari_listele()}
        self.eid=optik.sinav_ekle('Camera test','2026-10-17','camera-source.pdf',list('ABCDABCDABCDAB'),self.t['id'],'yaprak','Fen Bilimleri',self.student['sinif_id'])
        (optik.klasor()/'camera-source.pdf').write_bytes(pdf())
        con=kt.connect();con.execute('INSERT INTO kamera_sinav(sinav_id,choices) VALUES(?,?)',(self.eid,'ABCD'));con.commit();con.close()
        con=kt.connect();con.execute("UPDATE lgs_optik_sinav SET kaynak='kamera' WHERE id=?",(self.eid,));con.commit();con.close()
        self.exam=optik.sinav_getir(self.eid)

    def tearDown(self):
        con=kt.connect()
        for exam in optik.sinavlari_listele():
            if exam['id'] not in self.before_ids:
                for table in ('kamera_form','kamera_sinav','kamera_islem','lgs_optik_sonuc'):
                    con.execute(f'DELETE FROM {table} WHERE sinav_id=?',(exam['id'],))
                con.execute('DELETE FROM lgs_optik_sinav WHERE id=?',(exam['id'],))
        con.commit();con.close()

    def form(self,marks=None,extra=None):
        r=self.c.get(f'/yaprak-kamera/{self.eid}/yazdir?student={self.student["id"]}')
        self.assertEqual(r.status_code,200);self.assertEqual(len(PdfReader(BytesIO(r.data)).pages),2)
        con=kt.connect();token=con.execute('SELECT token FROM kamera_form WHERE sinav_id=?',(self.eid,)).fetchone()[0];con.close()
        im=marked_form(self.exam,self.student,token,marks or self.exam['cevaplar'],extra)
        return im

    def scan(self,im):
        _,jpg=cv2.imencode('.jpg',im)
        return self.c.post(f'/yaprak-kamera/{self.eid}/tara',data={'csrf':self.csrf,'image':(BytesIO(jpg.tobytes()),'scan.jpg')})

    def save(self,scan,**overrides):
        data=dict(csrf=self.csrf,ticket=scan['ticket'],answers=json.dumps(scan['answers']));data.update(overrides)
        return self.c.post(f'/yaprak-kamera/{self.eid}/kaydet',data=data)

    def test_prepare_confirm_ownership_and_print(self):
        r=self.c.post('/yaprak-kamera/hazirla',data=dict(csrf=self.csrf,count='4',choices='ABCD',class_id=str(self.student['sinif_id']),name='Uploaded test',course='Fen',questions=(BytesIO(pdf()),'test.pdf'),key=(BytesIO(pdf('1 A 2 B 3 C 4 D')),'key.pdf')))
        self.assertEqual(r.status_code,302);path=r.location
        draft_html=self.c.get(path).get_data(as_text=True)
        self.assertIn('Anahtarı onayla',draft_html)
        self.assertIn('4 / 4 cevap belgeden okundu',draft_html)
        preview=self.c.get(path+'/onizleme/1')
        self.assertEqual(preview.status_code,200);self.assertEqual(preview.mimetype,'image/png')
        self.assertEqual(self.c.get(path+'/onizleme/9').status_code,404)
        self.assertEqual(self.c.post(path,data=dict(csrf=self.csrf)).status_code,200)
        r=self.c.post(path,data=dict(csrf=self.csrf,q1='A',q2='B',q3='C',q4='D',confirmed='yes'))
        self.assertEqual(r.status_code,302);self.assertEqual(self.c.get(r.location).status_code,200)
        self.assertEqual(self.c.get(path).status_code,404)

    def test_clear_scan_save_duplicate_and_parent_analysis(self):
        r=self.scan(self.form());self.assertEqual(r.status_code,200,r.data)
        self.assertEqual(r.json['answers'],self.exam['cevaplar']);self.assertEqual(r.json['uncertain'],[])
        saved=self.save(r.json);self.assertEqual(saved.status_code,200,saved.data);self.assertEqual(saved.json['dogru'],14)
        self.assertEqual(self.save(r.json).status_code,409)
        again=self.scan(self.form());self.assertTrue(again.json['existing'])
        self.assertEqual(self.save(again.json).status_code,409)
        self.assertEqual(self.save(again.json,replace='yes',reviewed='yes').status_code,200)
        parent=w.app.test_client()
        import veli_auth
        con=veli_auth.connect()
        con.execute('INSERT OR REPLACE INTO veli_credentials(ogrenci_id,password_hash,auth_version) VALUES(?,?,?)',(self.student['id'],'test-fixture','camera-test-version'))
        con.commit();con.close()
        with parent.session_transaction() as s:
            s['veli_ogrenci_id']=self.student['id']
            s['veli_auth_version']='camera-test-version'
        response=parent.get(f'/evde-analiz?test={self.eid}')
        self.assertEqual(response.status_code,200)
        school=parent.get('/okul-yaprak-analiz').get_data(as_text=True)
        self.assertIn('Soru bazlı cevapları incele',school)
        self.assertEqual(parent.get('/yaprak-kamera').status_code,302)

    def test_perspective_rotation_and_double_mark(self):
        im=self.form(extra=1)
        height,width=im.shape[:2]
        src=np.float32([[0,0],[width-1,0],[width-1,height-1],[0,height-1]])
        dst=np.float32([[100,50],[width-120,0],[width-25,height-80],[20,height-20]])
        warped=cv2.warpPerspective(im,cv2.getPerspectiveTransform(src,dst),(width,height),borderValue=(255,255,255))
        warped=cv2.rotate(warped,cv2.ROTATE_90_CLOCKWISE)
        result=self.scan(warped);self.assertEqual(result.status_code,200,result.data)
        self.assertIn(1,result.json['uncertain'])
        self.assertEqual(self.save(result.json).status_code,400)
        answers=result.json['answers'];answers[0]='A'
        self.assertEqual(self.save(result.json,answers=json.dumps(answers),reviewed='yes').status_code,200)

    def test_blank_requires_review_crop_rejected(self):
        result=self.scan(self.form(marks=['']*14));self.assertEqual(result.status_code,200,result.data)
        self.assertEqual(len(result.json['uncertain']),14)
        self.assertEqual(self.scan(self.form()[150:]).status_code,400)

    def test_changed_key_rejects_pending_scan(self):
        result=self.scan(self.form());self.assertEqual(result.status_code,200,result.data)
        con=kt.connect();con.execute('UPDATE lgs_optik_sinav SET cevaplar=? WHERE id=?',(json.dumps(['D']*14),self.eid));con.commit();con.close()
        self.assertEqual(self.save(result.json).status_code,409)

    def test_permissions_csrf_and_wrong_test(self):
        self.assertEqual(self.c.post(f'/yaprak-kamera/{self.eid}/kaydet').status_code,400)
        other=w.app.test_client();teacher=next(t for t in d.tum_ogretmenler() if t['yetki']=='tam' and t['id']!=self.t['id'])
        with other.session_transaction() as s:s['ogretmen_id']=teacher['id']
        self.assertEqual(other.get(f'/yaprak-kamera/{self.eid}').status_code,404)
        im=self.form();con=kt.connect();con.execute('DELETE FROM kamera_form WHERE sinav_id=?',(self.eid,));con.commit();con.close()
        self.assertEqual(self.scan(im).status_code,400)


if __name__=='__main__':unittest.main()
