"""Document conversion, conservative answer-key extraction and printed-form OMR."""
from __future__ import annotations

from functools import lru_cache
from io import BytesIO
from pathlib import Path
import re
import os
import shutil
import subprocess
import tempfile
import threading
import zipfile

from PIL import Image, ImageOps
from pypdf import PdfReader, PdfWriter

MAX_PAGES = 30
SCALE = 6  # Canonical scan pixels per millimetre.
MARKERS = [(8, 8), (192, 8), (192, 279), (8, 279)]
OCR_LOCK = threading.Lock()


def image_open(data):
    im = Image.open(BytesIO(data))
    if im.width * im.height > 24_000_000:
        raise ValueError('Görsel çok büyük. 24 megapikselden küçük bir görsel yükleyin.')
    im = ImageOps.exif_transpose(im).convert('RGB')
    im.thumbnail((2600, 3600))
    return im


def valid_pdf(data):
    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted or not 1 <= len(reader.pages) <= MAX_PAGES:
        raise ValueError('PDF şifresiz ve 1-30 sayfa olmalı.')
    return reader


def docx_check(data):
    with zipfile.ZipFile(BytesIO(data)) as archive:
        if 'word/document.xml' not in archive.namelist():
            raise ValueError('Geçerli bir Word (.docx) belgesi seçin.')
        if sum(x.file_size for x in archive.infolist()) > 80 * 1024 * 1024:
            raise ValueError('Word belgesinin açılmış boyutu çok büyük.')
        if any('vbaProject' in x for x in archive.namelist()):
            raise ValueError('Makrolu Word belgesi kabul edilmez.')
        for name in archive.namelist():
            if name.endswith('.rels'):
                from lxml import etree
                root = etree.fromstring(archive.read(name), etree.XMLParser(resolve_entities=False, no_network=True))
                for rel in root:
                    if rel.get('TargetMode') == 'External' and not rel.get('Type', '').endswith('/hyperlink'):
                        raise ValueError('Word belgesindeki dış bağlantılı görselleri belgeye gömerek yeniden yükleyin.')


def convert_document(data, suffix):
    suffix = suffix.lower()
    if suffix == '.pdf':
        valid_pdf(data)
        return data
    if suffix in {'.jpg', '.jpeg', '.png'}:
        from reportlab.pdfgen import canvas
        from reportlab.lib.utils import ImageReader
        from reportlab.lib.pagesizes import A4
        im = image_open(data)
        out = BytesIO(); c = canvas.Canvas(out, pagesize=A4)
        w, h = A4
        ratio = min((w-36)/im.width, (h-36)/im.height)
        iw, ih = im.width*ratio, im.height*ratio
        c.drawImage(ImageReader(im), (w-iw)/2, (h-ih)/2, iw, ih)
        c.save()
        return out.getvalue()
    if suffix != '.docx':
        raise ValueError('PDF, Word (.docx), JPEG veya PNG seçin. Eski .doc dosyalarını .docx olarak kaydedin.')
    docx_check(data)
    from broadcast_presentations import office_command
    exe = office_command()
    if not exe:
        raise ValueError('Word dönüştürücü bu sunucuda hazır değil. Belgeyi PDF olarak yükleyebilirsiniz.')
    with tempfile.TemporaryDirectory(prefix='kamera-word-') as folder:
        folder = Path(folder)
        source = folder / 'source.docx'; source.write_bytes(data)
        profile=folder/'profile';(profile/'user').mkdir(parents=True)
        (profile/'user/registrymodifications.xcu').write_text('''<?xml version="1.0"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry">
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Load"><prop oor:name="UpdateLinks" oor:op="fuse"><value>false</value></prop></item>
</oor:items>''',encoding='utf-8')
        try:
            subprocess.run([exe, '-env:UserInstallation='+profile.as_uri(),
                            '--headless', '--nologo', '--nodefault', '--norestore', '--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(folder), str(source)],
                           timeout=50, check=True, capture_output=True,
                           creationflags=subprocess.CREATE_NO_WINDOW if os.name=='nt' else 0)
        except (subprocess.SubprocessError, OSError) as exc:
            raise ValueError('Word belgesi dönüştürülemedi. PDF olarak kaydedip tekrar deneyin.') from exc
        target = folder / 'source.pdf'
        if not target.exists():
            raise ValueError('Word belgesi için PDF oluşturulamadı.')
        result = target.read_bytes(); valid_pdf(result)
        return result


def merge_documents(documents):
    writer = PdfWriter()
    for data, suffix in documents:
        reader = valid_pdf(convert_document(data, suffix))
        if len(writer.pages) + len(reader.pages) > MAX_PAGES:
            raise ValueError('Bir test en fazla 30 soru sayfası içerebilir.')
        for page in reader.pages:
            writer.add_page(page, excluded_keys=['/AA', '/Annots'])
    out = BytesIO(); writer.write(out)
    return out.getvalue(), len(writer.pages)


@lru_cache(maxsize=1)
def ocr_engine():
    from rapidocr_onnxruntime import RapidOCR
    return RapidOCR(text_score=0.5, intra_op_num_threads=2, inter_op_num_threads=1,
                    max_side_len=1600, det_limit_type='max', det_limit_side_len=1600)


def ocr_image(im):
    import numpy as np
    from PIL import ImageChops
    # Read overlapping strips rather than a full high-resolution page tensor.
    # Keep the original glyph size while bounding inference memory on the host.
    rows = []
    with OCR_LOCK:
        for top in range(0, im.height, 750):
            strip = im.crop((0, top, im.width, min(top + 900, im.height))).convert('RGB')
            ink = ImageChops.difference(strip, Image.new('RGB', strip.size, 'white')).convert('L')
            bounds = ink.point(lambda value: 255 if value > 45 else 0).getbbox()
            if not bounds:
                continue
            left = max(0, bounds[0]-16); upper = max(0, bounds[1]-16)
            strip = strip.crop((left, upper, min(strip.width,bounds[2]+16), min(strip.height,bounds[3]+16)))
            detected, _ = ocr_engine()(np.array(strip)[:, :, ::-1])
            for box, text, confidence in detected or []:
                box = [[x+left, y+top+upper] for x,y in box]
                cy = sum(p[1] for p in box)/4; cx = sum(p[0] for p in box)/4
                if any(re.sub(r'\s+','',text) == re.sub(r'\s+','',old[1]) and abs(cy-sum(p[1] for p in old[0])/4)<12
                       and abs(cx-sum(p[0] for p in old[0])/4)<12 for old in rows):
                    continue
                rows.append((box,text,confidence))
    # Reconstruct spatial lines so a number and answer in separate OCR boxes stay together.
    lines = []
    for box, text, confidence in sorted(rows or [], key=lambda r: min(p[1] for p in r[0])):
        if confidence < .8:
            text = '?'  # Never silently turn a low-confidence recognition into an answer.
        cy = sum(p[1] for p in box)/4
        height = max(p[1] for p in box)-min(p[1] for p in box)
        line = next((r for r in lines if abs(r['y']-cy) < max(5, height*.45)), None)
        if line is None:
            line = {'y': cy, 'parts': []}; lines.append(line)
        line['parts'].append((min(p[0] for p in box), text))
    return '\n'.join(' '.join(t for _, t in sorted(line['parts'])) for line in lines)


def extract_key_text(data, suffix):
    if suffix == '.docx':
        docx_check(data)
        from docx import Document
        doc = Document(BytesIO(data))
        from docx.oxml.ns import qn
        lines = []
        for child in doc.element.body:
            if child.tag == qn('w:p'):
                lines.append(''.join(el.text or '' for el in child.iter(qn('w:t'))))
            elif child.tag == qn('w:tbl'):
                for row in child.findall(qn('w:tr')):
                    lines.append(' '.join(''.join(el.text or '' for el in cell.iter(qn('w:t'))) for cell in row.findall(qn('w:tc'))))
        return '\n'.join(lines), False
    if suffix in {'.jpg', '.jpeg', '.png'}:
        return ocr_image(image_open(data)), True
    if suffix != '.pdf':
        raise ValueError('Cevap anahtarı PDF, DOCX, JPEG veya PNG olmalı.')
    reader = valid_pdf(data)
    if len(reader.pages) > 5:
        raise ValueError('Cevap anahtarı en fazla 5 sayfa olmalı; ilgili sayfaları ayırıp yükleyin.')
    import pypdfium2 as pdfium
    result = []; used_ocr = False
    with pdfium.PdfDocument(data) as pdf:
        for i, page in enumerate(reader.pages):
            text = page.extract_text(extraction_mode='layout') or ''
            # A scanned key can still have a text footer/year. Digits alone do
            # not mean the actual answer table was extracted successfully.
            if not any(parse_key(text, 90, 'ABCDE')[0]):
                used_ocr = True
                rendered = pdf[i].render(scale=3)
                text = ocr_image(rendered.to_pil())
                rendered.close()
            result.append(text)
    return '\n'.join(result), used_ocr


def parse_key(text, count, choices):
    """Only numbered pairs or explicit two-row tables, never arbitrary letters in prose."""
    text = text.upper().replace('İ', 'I')
    for dash in ('–','—','−','‑'):
        text=text.replace(dash,'-')
    found = {}; conflicts = set(); warnings = []
    def add(number, answer):
        if not 1 <= number <= count:
            warnings.append(f'{number}. soru testin soru sayısı dışında.'); return
        if number in found:
            conflicts.add(number)
        found[number] = answer if answer in choices else ''
    lines = [r.strip() for r in text.splitlines() if r.strip()]
    # Table row labels are not answer content. Keep all other prose intact.
    lines=[re.sub(r'^(?:SORU(?:\s*(?:NO(?:SU)?|NUMARASI))?|CEVAP(?:LAR[I]?|\s*ANAHTARI)?)\s*[:|]?\s*(?=[\dA-E])','',line) for line in lines]
    consumed=set()
    # PDFs sometimes serialize a two-column table as one cell per line.
    for i,line in enumerate(lines):
        if i in consumed:continue
        if re.fullmatch(r'[1-9]\d?[.)\-:]?',line):
            j=i
            while j<len(lines) and re.fullmatch(r'[1-9]\d?[.)\-:]?',lines[j]):j+=1
            n=j-i
            answers=lines[j:j+n]
            if len(answers)==n and all(re.fullmatch(r'[A-E]',a) for a in answers):
                for number,answer in zip(lines[i:j],answers):add(int(re.search(r'\d+',number)[0]),answer)
                consumed.update(range(i,j+n))
    for i, line in enumerate(lines):
        if i in consumed:continue
        pairs=list(re.finditer(r'(?<![\w])([1-9]\d?)\s*[.\-:)=|]?\s*\(?([A-E])\)?(?![\w])', line))
        for j,pair in enumerate(pairs):
            tail=line[pair.end():pairs[j+1].start() if j+1<len(pairs) else len(line)]
            ambiguous=bool(re.match(r'\s*[/,;(]?\s*[A-E]\b',tail))
            add(int(pair[1]), '' if ambiguous else pair[2])
        nums = re.findall(r'\d+', line)
        if nums and re.fullmatch(r'[\d\s|,;.]+', line) and i+1 < len(lines):
            answers = re.findall(r'\b[A-E]\b', lines[i+1])
            if len(answers) == len(nums) and re.fullmatch(r'[A-E\s|,;.]+', lines[i+1]):
                for n, a in zip(nums, answers): add(int(n), a)
    for number in conflicts: found[number] = ''
    if conflicts: warnings.append('Tekrarlanan soru numaraları var. Kitapçıkları ayrı test olarak yükleyin: '+', '.join(map(str, sorted(conflicts))))
    answers = [found.get(i, '') for i in range(1, count+1)]
    missing = [str(i+1) for i,a in enumerate(answers) if not a]
    if missing: warnings.append('Kontrol edilmesi gereken sorular: '+', '.join(missing))
    return answers, list(dict.fromkeys(warnings))


def bubble_position(index, option, count=90, version=1):
    if version==2:
        rows=20 if count<=60 else 30
        return 30+(index//rows)*62+option*8, 72+(index%rows)*(9 if rows==20 else 6.2)
    return 28+(index//30)*62+option*8, 68+(index%30)*6.2


def form_pdf(exam, student, token, choices, version=2):
    import cv2
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.graphics.barcode import qr
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics import renderPDF
    font = Path(__file__).parent/'static/vendor/pdfjs/standard_fonts/LiberationSans-Regular.ttf'
    if 'CameraSans' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('CameraSans', str(font)))
    out = BytesIO(); c = canvas.Canvas(out, pagesize=(210*mm,297*mm))
    def text(x,y,value,size=10):
        c.setFont('CameraSans',size); c.drawString(x*mm,(297-y)*mm,str(value))
    text(23,17,'AKADEMİPUAN | YAPRAK TEST',12)
    name=exam['ad']
    while pdfmetrics.stringWidth(name,'CameraSans',10)>128*mm: name=name[:-2]+'…'
    text(23,26,name)
    text(23,35,student['ad_soyad'][:48],11)
    text(23,43,f"{student['sinif_adi']} / No: {student['ogr_no']} / {exam['soru_sayisi']} soru",9)
    text(23,51,'Her soruda tek yuvarlağı tamamen doldurun.',9)
    text(23,57,'Değiştirirken eski işareti tamamen silin.',9)
    code=qr.QrCodeWidget(('AP2:' if version==2 else 'AP1:')+token); box=code.getBounds()
    drawing=Drawing(33*mm,33*mm,transform=[33*mm/(box[2]-box[0]),0,0,33*mm/(box[3]-box[1]),0,0])
    drawing.add(code);renderPDF.draw(drawing,c,161*mm,(297-55)*mm)
    dictionary=cv2.aruco.getPredefinedDictionary(cv2.aruco.DICT_4X4_50)
    for i,(x,y) in enumerate(MARKERS):
        marker=cv2.aruco.generateImageMarker(dictionary,i,200)
        size=12 if version==2 else 10
        c.drawImage(ImageReader(Image.fromarray(marker)),(x+5-size/2)*mm,(297-y-5-size/2)*mm,size*mm,size*mm)
    rows=20 if version==2 and exam['soru_sayisi']<=60 else 30
    for col in range((exam['soru_sayisi']+rows-1)//rows):
        for option,letter in enumerate(choices):
            text((29 if version==2 else 27)+col*62+option*8,65 if version==2 else 63,letter,9)
    for i in range(exam['soru_sayisi']):
        x,y=bubble_position(i,0,exam['soru_sayisi'],version);text(x-9,y+1,str(i+1),9)
        for j in range(len(choices)):
            x,y=bubble_position(i,j,exam['soru_sayisi'],version)
            c.setLineWidth(.6);c.circle(x*mm,(297-y)*mm,(2.5 if version==2 else 2)*mm,stroke=1,fill=0)
    text(23,267,'Bu sayfayı testle birlikte teslim edin. Katlamayın ve QR alanını karalamayın.',8)
    text(23,274,'Öğretmen: Optikleri tek PDF olarak tarayın ve Toplu PDF bölümüne yükleyin.',8)
    text(23,287,('OPTİK 2 / ' if version==2 else '')+'A4 / Gerçek boyut (%100) / Tek yüz baskı',8)
    c.save();return out.getvalue()


def read_form(data, count, choices):
    import cv2
    import numpy as np
    im=image_open(data); frame=cv2.cvtColor(np.array(im),cv2.COLOR_RGB2GRAY)
    detector=cv2.aruco.ArucoDetector(cv2.aruco.getPredefinedDictionary(cv2.aruco.DICT_4X4_50))
    # Search a smaller image first; keep the original for answer sampling.
    ratio=min(1.0,1800/max(frame.shape))
    detection=cv2.resize(frame,None,fx=ratio,fy=ratio,interpolation=cv2.INTER_AREA) if ratio<1 else frame
    corners,ids,_=detector.detectMarkers(detection)
    if ratio<1 and ids is not None and all(list(ids.flatten()).count(i)==1 for i in range(4)):
        corners=[corner/ratio for corner in corners]
    elif ratio<1:
        corners,ids,_=detector.detectMarkers(frame)
    if ids is None or any(list(ids.flatten()).count(i)!=1 for i in range(4)):
        raise ValueError('Dört köşe işareti okunamadı. Sayfanın tamamını gösterin; ışığı ve netliği kontrol edin.')
    locations={int(i):corner[0].mean(axis=0) for i,corner in zip(ids.flatten(),corners)}
    source=np.float32([locations[i] for i in range(4)])
    if abs(cv2.contourArea(source)) < 140_000:
        raise ValueError('Sayfa görüntüde çok küçük. Kamerayı biraz yaklaştırın.')
    target=np.float32([[(x+5)*SCALE,(y+5)*SCALE] for x,y in MARKERS])
    warp=cv2.warpPerspective(frame,cv2.getPerspectiveTransform(source,target),(210*SCALE,297*SCALE),borderValue=255)
    # QR is read only after geometric normalization; the signed database token identifies the pupil.
    import zxingcpp
    decoded=zxingcpp.read_barcode(warp[18*SCALE:58*SCALE,157*SCALE:198*SCALE],formats=zxingcpp.BarcodeFormat.QRCode)
    code=decoded.text if decoded else ''
    if not code.startswith(('AP1:','AP2:')): raise ValueError('Öğrenci QR kodu okunamadı. QR alanını netleştirip tekrar çekin.')
    if cv2.Laplacian(warp,cv2.CV_64F).var()<30:
        raise ValueError('Görüntü bulanık. Telefonu sabitleyip yeniden çekin.')
    version=2 if code.startswith('AP2:') else 1
    binary=cv2.adaptiveThreshold(warp,255,cv2.ADAPTIVE_THRESH_GAUSSIAN_C,cv2.THRESH_BINARY_INV,41,12)
    answers=[];uncertain=[];scores=[]
    for i in range(count):
        values=[]
        for j in range(len(choices)):
            x,y=bubble_position(i,j,count,version); x=round(x*SCALE);y=round(y*SCALE);r=round(1.25*SCALE)
            patch=binary[y-r:y+r+1,x-r:x+r+1]
            yy,xx=np.ogrid[-r:r+1,-r:r+1];mask=xx*xx+yy*yy<=r*r
            values.append(float(np.mean(patch[mask]))/255)
        ordered=sorted(values,reverse=True)
        mark=choices[int(np.argmax(values))] if ordered[0]>.48 and ordered[1]<.18 else ''
        # A truly empty row is accepted; weak and multiple marks require a human decision.
        if not mark and ordered[0]>.07: uncertain.append(i+1)
        answers.append(mark);scores.append([round(v,3) for v in values])
    if not any(answers): uncertain=list(range(1,count+1))
    preview=cv2.resize(warp,(630,891));_,jpg=cv2.imencode('.jpg',preview,[cv2.IMWRITE_JPEG_QUALITY,80])
    return {'token':code[4:],'answers':answers,'uncertain':uncertain,'scores':scores,'preview':jpg.tobytes()}
